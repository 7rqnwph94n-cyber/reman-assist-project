"""
Costing / office sheet scaffolding.

These helpers mirror the BOM→Costing flow described by product, but are
intentionally minimal so they can be safely wired in later without breaking
current survey/BOM flows.
"""

import logging
import os
from typing import Any, Dict, Iterable, Tuple

from flask import current_app

from models import (
    db,
    Job,
    PartRequirement,
    CostingSheet,
    CostingTemplateProfile,
    CostingLine,
    LineClassEnum,
    MatchStatusEnum,
    ProcurementStatusEnum,
    CostingTemplateLine,
    JobCostingUpload,
)

log = logging.getLogger(__name__)


def ensure_costing_sheet(
    job: Job,
    profile: CostingTemplateProfile | None = None,
    upload: JobCostingUpload | None = None,
) -> CostingSheet:
    """
    Ensure a CostingSheet exists for the given job.
    - If one exists, return it (and update profile/upload pointers if supplied).
    - Otherwise, create it with the supplied profile/upload.
    """
    sheet = CostingSheet.query.filter_by(job_id=job.id).first()
    if sheet:
        if upload and sheet.job_costing_upload_id != upload.id:
            sheet.job_costing_upload_id = upload.id
        if profile and not sheet.costing_template_profile_id:
            sheet.costing_template_profile_id = profile.id
            sheet.template_version_snapshot = sheet.template_version_snapshot or profile.version
        db.session.commit()
        return sheet

    sheet = CostingSheet(
        job_id=job.id,
        job_costing_upload_id=upload.id if upload else None,
        costing_template_profile_id=profile.id if profile else None,
        template_version_snapshot=profile.version if profile else None,
        job_no=job.job_number,
        sn=job.serial_number,
        engine_desc=job.engine_type,
    )
    db.session.add(sheet)
    db.session.commit()
    log.info("Created costing sheet for job %s (profile=%s)", job.id, profile.id if profile else None)
    return sheet


def ensure_costing_sheet_from_template(
    job: Job, profile: CostingTemplateProfile | None = None
) -> CostingSheet:
    """
    Backwards compatible helper for template-based costing flows.
    """
    return ensure_costing_sheet(job, profile=profile, upload=None)


def create_costing_lines_from_bom(job: Job) -> int:
    """
    If no costing lines exist, seed from the current BOM (PartRequirement).
    """
    sheet = CostingSheet.query.filter_by(job_id=job.id).first()
    if not sheet:
        return 0
    existing = sheet.lines
    if existing:
        return 0
    parts = PartRequirement.query.filter_by(job_id=job.id).all()
    created = 0
    for pr in parts:
        cl = CostingLine(
            sheet=sheet,
            part_requirement_id=pr.id,
            line_no=created + 1,
            group_no=pr.section,
            part_description=pr.name,
            system_part_number=pr.part_no,
            qty_required=pr.qty_required or 0,
            qty_on_stock=0,
            qty_on_order=0,
            intended_supplier=pr.supplier,
            po_number=pr.po_number,
            notes=pr.notes,
            match_status=MatchStatusEnum.MATCHED,
        )
        db.session.add(cl)
        created += 1
    db.session.commit()
    return created


def create_costing_lines_from_template(
    job: Job,
    profile: CostingTemplateProfile,
    template_rows: Iterable[Dict[str, Any]],
) -> int:
    """
    Seed costing lines for a job from parsed template rows.
    Returns number of lines created.
    """
    sheet = ensure_costing_sheet(job, profile=profile)
    created = 0

    for idx, row in enumerate(template_rows, start=1):
        cl = CostingLine(
            sheet=sheet,
            line_no=row.get("line_no") or idx,
            group_no=row.get("group_no"),
            figure=row.get("figure"),
            item_no=row.get("item_no"),
            oem_part_no=row.get("oem_part_no"),
            system_part_number=row.get("system_part_number"),
            part_description=row.get("part_description"),
            line_class=row.get("line_class") or LineClassEnum.STANDARD,
            is_menu=bool(row.get("is_menu", False)),
            is_kit=bool(row.get("is_kit", False)),
            is_consumable=bool(row.get("is_consumable", False)),
            intended_supplier=row.get("intended_supplier"),
            latest_supplier=row.get("latest_supplier"),
            latest_cost=row.get("latest_cost"),
            menu_sell_price_unit=row.get("menu_sell_price_unit"),
            qty_on_stock=row.get("qty_on_stock") or 0,
            qty_on_order=row.get("qty_on_order") or 0,
            qty_required=row.get("qty_required") or 0,
            procurement_status=ProcurementStatusEnum.NOT_COSTED,
            match_status=MatchStatusEnum.UNMATCHED_SHEET,
        )
        db.session.add(cl)
        created += 1

    db.session.commit()
    log.info("Seeded %s costing lines for job %s from template %s", created, job.id, profile.id)
    return created


def parse_costing_template_xlsx(profile: CostingTemplateProfile, file_path: str) -> Tuple[Dict[str, Any], list[Dict[str, Any]]]:
    """
    Placeholder parser for costing templates.
    Returns (header_values, rows) where rows are dicts matching CostingLine fields.
    """
    log.info("Parsing costing template %s from %s", profile.id, file_path)
    # Stub: hook in openpyxl parsing later
    return {}, []


def bind_costing_lines_to_bom(job: Job) -> int:
    """
    Attempt to bind existing costing lines to PartRequirement rows for a job.
    Returns number of matched lines.
    """
    sheet = CostingSheet.query.filter_by(job_id=job.id).first()
    if not sheet:
        return 0

    parts = PartRequirement.query.filter_by(job_id=job.id).all()
    pr_by_part = {(p.part_no or "").strip().lower(): p for p in parts if p.part_no}

    matched = 0
    for line in sheet.lines:
        keys = [
            (line.system_part_number or "").strip().lower(),
            (line.oem_part_no or "").strip().lower(),
        ]
        keys = [k for k in keys if k]
        found = []
        for k in keys:
            pr = pr_by_part.get(k)
            if pr:
                found.append(pr)
        if len(found) == 1:
            pr = found[0]
            line.part_requirement = pr
            line.qty_required = pr.qty_required or line.qty_required or 0
            line.match_status = MatchStatusEnum.MATCHED
            matched += 1
        elif len(found) > 1:
            line.match_status = MatchStatusEnum.AMBIGUOUS
        else:
            line.match_status = MatchStatusEnum.UNMATCHED_SHEET
    db.session.commit()
    return matched


def enrich_bom_from_catalogue(job: Job, preserve_overrides: bool = True) -> Dict[str, int]:
    """
    Enrich costing lines from template catalogue defaults, preserving overrides if requested.
    """
    sheet = CostingSheet.query.filter_by(job_id=job.id).first()
    if not sheet:
        return {"enriched": 0, "missing": 0}

    profile = sheet.template_profile
    if not profile:
        return {"enriched": 0, "missing": len(sheet.lines)}

    tmpl_lines = (
        CostingTemplateLine.query.filter_by(costing_template_profile_id=profile.id)
        .order_by(CostingTemplateLine.line_no, CostingTemplateLine.id)
        .all()
    )
    map_by_rec = {}
    map_by_oem = {}
    for tl in tmpl_lines:
        if tl.system_part_number:
            map_by_rec[tl.system_part_number.strip().lower()] = tl
        if tl.oem_part_no:
            map_by_oem[tl.oem_part_no.strip().lower()] = tl

    enriched = 0
    missing = 0
    for line in sheet.lines:
        key_rec = (line.system_part_number or "").strip().lower()
        key_oem = (line.oem_part_no or "").strip().lower()
        tl = None
        if key_rec and key_rec in map_by_rec:
            tl = map_by_rec[key_rec]
        elif key_oem and key_oem in map_by_oem:
            tl = map_by_oem[key_oem]

        if not tl:
            missing += 1
            continue

        def set_if_empty(attr, value):
            if not preserve_overrides or getattr(line, attr) in (None, "", 0):
                setattr(line, attr, value)

        set_if_empty("line_class", tl.default_class or LineClassEnum.STANDARD)
        set_if_empty("intended_supplier", tl.default_intended_supplier)
        set_if_empty("latest_supplier", tl.default_latest_supplier)
        set_if_empty("latest_cost", tl.default_latest_cost)
        set_if_empty("menu_sell_price_unit", tl.default_menu_sell_price)
        set_if_empty("qty_on_stock", tl.default_qty_on_stock)
        set_if_empty("qty_on_order", tl.default_qty_on_order)
        # flags always sync unless overrides explicitly set
        if not preserve_overrides:
            line.is_menu = tl.is_menu
            line.is_kit = tl.is_kit
            line.is_consumable = tl.is_consumable
        else:
            line.is_menu = line.is_menu or tl.is_menu
            line.is_kit = line.is_kit or tl.is_kit
            line.is_consumable = line.is_consumable or tl.is_consumable

        enriched += 1

    db.session.commit()
    return {"enriched": enriched, "missing": missing}


def export_costing_sheet_xlsx(job: Job, output_path: str) -> None:
    """
    Export using stored Excel template, preserving formatting.
    """
    sheet = CostingSheet.query.filter_by(job_id=job.id).first()
    if not sheet:
        raise ValueError("No costing sheet for job")
    profile = sheet.template_profile
    if not profile or not profile.xlsx_template_path:
        raise ValueError("No costing template file configured")

    from openpyxl import load_workbook

    template_path = profile.xlsx_template_path
    if template_path and not os.path.isabs(template_path):
        template_path = os.path.join(current_app.config.get("UPLOAD_FOLDER", ""), template_path)
    wb = load_workbook(template_path)
    ws = wb[profile.sheet_name] if profile.sheet_name and profile.sheet_name in wb.sheetnames else wb.active

    # Fill header cells
    header_map = profile.header_cell_map_json or {}
    header_values = {
        "job_no": sheet.job_no or job.job_number,
        "sn": sheet.sn or job.serial_number,
        "engine_no": sheet.engine_no or job.engine_type,
        "engine_desc": sheet.engine_desc or job.engine_type,
        "issue": sheet.issue,
        "auth": sheet.auth,
        "date": sheet.date,
        "parts_replacement_cost": sheet.parts_replacement_cost,
        "emergent_work_cost": sheet.emergent_work_cost,
        "menu_price_add": sheet.menu_price_add,
        "cost_booked": sheet.cost_booked,
        "warning_text": sheet.warning_text,
        "stage": sheet.stage.value if sheet.stage else None,
    }
    for field, cell_addr in header_map.items():
        if not cell_addr:
            continue
        ws[cell_addr] = header_values.get(field)

    start_row = profile.table_start_row or (ws.min_row + 1)
    header_row_idx = start_row - 1 if start_row > 1 else start_row
    header_cells = {cell.value: cell.column for cell in ws[header_row_idx] if cell.value}

    logical_to_field = {
        "line_no": "line_no",
        "group_no": "group_no",
        "figure": "figure",
        "item_no": "item_no",
        "system_part_number": "system_part_number",
        "oem_part_no": "oem_part_no",
        "part_description": "part_description",
        "notes": "notes",
        "line_class": "line_class",
        "intended_supplier": "intended_supplier",
        "latest_supplier": "latest_supplier",
        "latest_cost": "latest_cost",
        "price_cost": "price_cost",
        "menu_sell_price_unit": "menu_sell_price_unit",
        "qty_required": "qty_required",
        "qty_required_dup": "qty_required",
        "qty_on_stock": "qty_on_stock",
        "qty_on_order": "qty_on_order",
        "po_number": "po_number",
        "unit_cost_confirmed": "unit_cost_confirmed",
        "line_cost": "line_cost",
        "procurement_status": "procurement_status",
        "emergent_line_cost": "emergent_line_cost",
        "emergent_qty_required": "emergent_qty_required",
        "rbb_value": "rbb_value",
        "is_menu": "is_menu",
        "is_kit": "is_kit",
        "is_consumable": "is_consumable",
    }

    col_map = {}
    req_map = profile.required_columns_json or {}
    for logical, header_text in req_map.items():
        for hdr, col in header_cells.items():
            if str(hdr).strip().lower() == str(header_text).strip().lower():
                col_map[logical] = col
                break

    def write_cell(row_idx: int, col_identifier, value):
        try:
            ws.cell(row=row_idx, column=col_identifier, value=value)
        except Exception:
            pass

    # Use template row height as baseline if present
    default_row_height = None
    try:
        default_row_height = ws.row_dimensions[start_row].height or ws.row_dimensions[start_row - 1].height
    except Exception:
        default_row_height = ws.row_dimensions[ws.min_row].height if ws.row_dimensions.get(ws.min_row) else None
    if not default_row_height:
        default_row_height = 18  # sensible default

    for idx, line in enumerate(sheet.lines, start=start_row):
        # Keep rows visually consistent with template height
        ws.row_dimensions[idx].height = default_row_height
        for logical, field in logical_to_field.items():
            if logical not in col_map:
                continue
            val = getattr(line, field, None)
            if logical == "part_description" and not val and getattr(line, "part_requirement", None):
                val = getattr(line.part_requirement, "name", None)
            if logical == "procurement_status" and val is not None:
                val = val.value if hasattr(val, "value") else val
            if logical in {"line_class"} and val is not None and hasattr(val, "value"):
                val = val.value
            if logical in {"is_menu", "is_kit", "is_consumable"}:
                val = 1 if val else 0
            write_cell(idx, col_map[logical], val)

    wb.save(output_path)
    log.info("Exported costing sheet for job %s to %s", job.id, output_path)
